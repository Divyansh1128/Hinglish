import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "faculty_synopsis.md"
OUTPUT = ROOT / "Hinglish_Math_Project_Synopsis.docx"


def add_runs(paragraph, text):
    for index, part in enumerate(re.split(r"\*\*(.*?)\*\*", text)):
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1


def set_cell_padding(cell, amount=85):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        margin = OxmlElement(f"w:{edge}")
        margin.set(qn("w:w"), str(amount))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 21), ("Subtitle", 12), ("Heading 1", 12)):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(9 if name == "Heading 1" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    document.styles["List Bullet"].font.name = "Times New Roman"
    document.styles["List Bullet"].font.size = Pt(11)
    document.styles["List Bullet"].paragraph_format.space_after = Pt(5)
    document.styles["List Bullet"].paragraph_format.line_spacing = 1.08
    for style in document.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "List Bullet"):
        fonts = document.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn(f"w:{attribute}"), None)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attribute}"), "Times New Roman")
    document.core_properties.title = "Hinglish Mathematical Reasoning in Large Language Models"
    document.core_properties.subject = "Bachelor of Technology major project synopsis for faculty review"
    document.core_properties.author = ""
    document.core_properties.keywords = "Hinglish, mathematical reasoning, code mixing, controlled evaluation"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Major Project Synopsis  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    lines = SOURCE.read_text().splitlines()
    position = 0
    while position < len(lines):
        line = lines[position].strip()
        position += 1
        if not line:
            continue
        if line == "<!-- pagebreak -->":
            document.add_page_break()
        elif line.startswith("### "):
            document.add_paragraph(line[4:], "Heading 1")
        elif line.startswith("## "):
            document.add_paragraph(line[3:], "Subtitle")
        elif line.startswith("# "):
            paragraph = document.add_paragraph(line[2:], "Title")
            for border in paragraph._p.xpath(".//w:pBdr"):
                border.getparent().remove(border)
        elif line.startswith("| "):
            rows = [line]
            while position < len(lines) and lines[position].startswith("| "):
                rows.append(lines[position])
                position += 1
            table = document.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            table.columns[0].width = Inches(1.2)
            table.columns[1].width = Inches(5.6)
            table.style = "Table Grid"
            data = [row.strip("|").split("|") for row in rows if not re.match(r"\|\s*---", row)]
            for row_index, values in enumerate(data):
                cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
                row_properties = table.rows[row_index]._tr.get_or_add_trPr()
                row_properties.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    row_properties.append(OxmlElement("w:tblHeader"))
                for cell_index, value in enumerate(values):
                    cell = cells[cell_index]
                    cell.width = Inches(1.2 if cell_index == 0 else 5.6)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    set_cell_padding(cell)
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    run = paragraph.add_run(value.strip())
                    run.font.size = Pt(10)
                    run.bold = row_index == 0
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "EEEEEE")
                        cell._tc.get_or_add_tcPr().append(shading)
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)
            spacer.paragraph_format.line_spacing = Pt(3)
            spacer.add_run().font.size = Pt(3)
        elif line.startswith("- "):
            add_runs(document.add_paragraph(style="List Bullet"), line[2:])
        else:
            paragraph = document.add_paragraph()
            add_runs(paragraph, line)
            if re.match(r"\[\d+\]", line):
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            if line.startswith(("Student name:", "Department:", "Date:")):
                paragraph.paragraph_format.space_after = Pt(5)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
